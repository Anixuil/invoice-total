"""Runs in a separate, single worker process; no invoice/statistics heuristics."""

import os
from io import BytesIO
from pathlib import Path
import zipfile
from .progress import report_progress

MAX_TEXT = 100000
_ocr = None


def parse_sources(files, database=None, job_id=None):
    os.environ.setdefault("OMP_NUM_THREADS", "1")
    parts = []
    characters = 0
    pages = 0
    failures = []

    def append(name, location, text, confidence=None):
        nonlocal characters
        text = str(text).strip()
        if not text:
            return
        characters += len(text)
        if characters > MAX_TEXT:
            raise ValueError("提取文字超过 10 万字符，请拆分材料")
        parts.append({"source": name, "location": location, "text": text,
                      "confidence": confidence, "needs_review": confidence is not None and confidence < 0.85})

    for file_index, file in enumerate(files):
        path = Path(file["path"])
        name = file["name"]
        suffix = path.suffix.lower()
        def progress(detail, completed=0, total=0):
            report_progress(database, job_id, f"文件 {file_index + 1}/{len(files)} · {name} · {detail}", completed, total)

        progress("正在读取文件")
        # Office zip containers are read in place, never extracted.
        if suffix in {".docx", ".xlsx"}:
            with zipfile.ZipFile(path) as archive:
                entries = archive.infolist()
                if len(entries) > 10000 or sum(item.file_size for item in entries) > 200 * 1024 * 1024:
                    raise ValueError(f"{name}：Office 文件展开体积超过限制")
        try:
            if suffix in {".txt", ".md"}:
                raw = path.read_bytes()
                try:
                    content = raw.decode("utf-8-sig")
                except UnicodeDecodeError:
                    content = raw.decode("gb18030")
                append(name, "全文", content)
            elif suffix == ".docx":
                from docx import Document
                from docx.oxml.ns import qn
                document = Document(path)
                group = []
                group_location = "正文"
                scores = []
                image_count = 0
                image_total = sum(node.tag in {qn("a:blip"), "{urn:schemas-microsoft-com:vml}imagedata"}
                                  for node in document.element.body.iter())

                def flush_group():
                    if group:
                        append(name, group_location, "\n".join(group), min(scores) if scores else None)
                    group.clear()
                    scores.clear()

                for index, child in enumerate(document.element.body):
                    if child.tag.endswith("}sectPr"):
                        continue
                    paragraphs = [child] if child.tag == qn("w:p") else list(child.iter(qn("w:p")))
                    for paragraph in paragraphs:
                        text = "".join(node.text or "" for node in paragraph.iter(qn("w:t"))).strip()
                        if text:
                            # A new textual paragraph starts a problem/context block; following
                            # image-only paragraphs stay attached to that block in document order.
                            flush_group()
                            group_location = f"正文块 {index + 1}"
                            group.append(text)
                        for node in paragraph.iter():
                            if node.tag not in {qn("a:blip"), "{urn:schemas-microsoft-com:vml}imagedata"}:
                                continue
                            image_count += 1
                            if image_count > 100:
                                raise ValueError("Word 内嵌图片超过 100 张，请拆分文档")
                            reference = node.get(qn("r:embed")) or node.get(qn("r:id"))
                            label = f"截图 {image_count}（正文块 {index + 1}）"
                            progress(f"正在识别截图 {image_count}/{image_total}", image_count - 1, image_total)
                            try:
                                part = document.part.related_parts.get(reference)
                                if part is None or not part.content_type.startswith("image/"):
                                    raise LookupError("Missing embedded image")
                                from PIL import Image
                                import numpy as np
                                with Image.open(BytesIO(part.blob)) as source:
                                    if source.width * source.height > 40_000_000:
                                        raise ValueError("截图像素过大")
                                    source.thumbnail((2600, 2600))
                                    array = np.asarray(source.convert("RGB"))[:, :, ::-1].copy()
                                recognized, confidence = ocr_image(array)
                                if not recognized.strip():
                                    raise LookupError("No readable text")
                            except Exception:
                                group.append(f"[{label}：未识别到可用文字，需人工查看原图，不得推断图中事实]")
                                failures.append(f"{name} / {label}：图片识别未完成，请人工核对原图")
                                scores.append(0)
                            else:
                                group.append(f"[{label} OCR，置信度 {confidence:.2f}]\n{recognized}")
                                scores.append(confidence)
                            progress(f"已处理截图 {image_count}/{image_total}（含识别失败项）", image_count, image_total)
                flush_group()
            elif suffix == ".xlsx":
                from openpyxl import load_workbook
                workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
                try:
                    cells = 0
                    for sheet_index, sheet in enumerate(workbook):
                        progress(f"正在读取工作表 {sheet.title}", sheet_index, len(workbook.sheetnames))
                        if sheet.max_row and sheet.max_column and sheet.max_row * sheet.max_column > 100000:
                            raise ValueError("工作表范围超过 10 万个单元格，请删除多余空白行列")
                        for index, row in enumerate(sheet.iter_rows(values_only=True), 1):
                            cells += len(row)
                            if cells > 100000:
                                raise ValueError("工作簿超过 10 万个单元格")
                            append(name, f"{sheet.title} / 行 {index}", " | ".join("" if v is None else str(v) for v in row))
                finally:
                    workbook.close()
            elif suffix == ".pdf":
                import pymupdf
                with pymupdf.open(path) as document:
                    if document.needs_pass:
                        raise ValueError("不支持加密 PDF")
                    pages += len(document)
                    if pages > 100:
                        raise ValueError("PDF 总页数超过 100 页")
                    for index, page in enumerate(document):
                        progress(f"正在识别第 {index + 1}/{len(document)} 页", index, len(document))
                        text = page.get_text().strip()
                        confidence = None
                        if len(text) < 20:
                            scale = min(2, 2200 / max(page.rect.width, page.rect.height))
                            pix = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False)
                            import numpy as np
                            array = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
                            text, confidence = ocr_image(array[:, :, :3][:, :, ::-1].copy())
                        append(name, f"第 {index + 1} 页", text, confidence)
                        progress(f"已处理 {index + 1}/{len(document)} 页", index + 1, len(document))
            else:
                from PIL import Image
                import numpy as np
                with Image.open(path) as source:
                    if source.width * source.height > 40_000_000:
                        raise ValueError("图片超过 4000 万像素")
                    source.thumbnail((2400, 2400))
                    image = np.asarray(source.convert("RGB"))[:, :, ::-1].copy()
                text, confidence = ocr_image(image)
                append(name, "图片 OCR", text, confidence)
        except ValueError:
            raise
        except Exception:
            failures.append(f"{name}：无法解析，请检查文件格式或改用粘贴文本")
        progress(f"文件处理完成（{file_index + 1}/{len(files)}）", file_index + 1, len(files))
    if not parts and not failures:
        failures.append("未识别到可用文字，请更换清晰材料或手工输入")
    return {"parts": parts, "errors": failures}


def ocr_image(image):
    global _ocr
    if _ocr is None:
        from rapidocr_onnxruntime import RapidOCR
        _ocr = RapidOCR(intra_op_num_threads=1, inter_op_num_threads=1)
    result, _ = _ocr(image)
    if not result:
        return "", 0
    return "\n".join(line[1] for line in result), min(float(line[2]) for line in result)
