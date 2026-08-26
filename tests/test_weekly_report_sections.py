from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches

from weekly_report_processor import (
    PPT_TEMPLATE,
    _flatten_shapes,
    _prefer_later_section_content,
    _section_from_entries,
    build_weekly_meeting_document,
    build_weekly_presentation,
    process_weekly_report,
)


def test_next_week_duplicate_is_not_kept_in_current_week() -> None:
    values = {
        "current": "已完成联调\n电子文档中心系统【并行测试】\n用户权限中心系统【并行测试】",
        "next": "电子文档中心系统【并行测试】\n用户权限中心系统【并行测试】",
        "issues": "无",
    }

    result = _prefer_later_section_content(values)

    assert result["current"] == "已完成联调"
    assert result["next"] == values["next"]
    assert result["issues"] == "无"


def test_issue_section_uses_the_right_column_boundary() -> None:
    entries = [
        {"text": "本周工作完成情况", "top": 994410, "left": 472440},
        {"text": "本周问题", "top": 990888, "left": 7780146},
        {"text": "下周计划", "top": 4932680, "left": 469900},
        {"text": "数据库连接异常", "top": 994410, "left": 7768590},
    ]

    assert _section_from_entries(entries, "issues") == "数据库连接异常"


def test_untemplated_project_ppt_is_added_to_weekly_presentation(tmp_path: Path) -> None:
    template = Presentation(PPT_TEMPLATE)
    presentation = Presentation()
    presentation.slide_width = template.slide_width
    presentation.slide_height = template.slide_height
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])

    for text, left, top, width, height in (
        ("新增自动识别项目（汇报人：张三）", 1, 0.3, 8, 0.4),
        ("本周工作完成情况", 0.4, 1.0, 1.5, 0.4),
        ("已完成联调", 1.2, 1.0, 5, 1),
        ("下周计划", 0.4, 2.2, 1.5, 0.4),
        ("安排上线", 1.2, 2.2, 5, 1),
        ("本周问题", 0.4, 3.4, 1.5, 0.4),
        ("风险项", 1.2, 3.4, 5, 1),
    ):
        slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height)).text = text

    source = tmp_path / "新增项目.pptx"
    presentation.save(source)
    result = process_weekly_report([(source, source.name)], [])

    added = next(item for item in result["projects"] if item["title"] == "新增自动识别项目")
    assert added["source_kind"] == "自动添加项目"
    assert added["current"] == "已完成联调"

    presentation_path = tmp_path / "项目周报.pptx"
    build_weekly_presentation(result, {source.name: source}, presentation_path)
    output = Presentation(presentation_path)
    slide = next(
        item for item in output.slides
        if "新增自动识别项目" in "".join(shape.text for shape in item.shapes if getattr(shape, "has_text_frame", False))
    )
    issue_body = next(entry for entry in _flatten_shapes(slide.shapes) if "风险项" in entry["text"])
    assert all(str(run.font.color.rgb) == "FF0000" for paragraph in issue_body["shape"].text_frame.paragraphs for run in paragraph.runs)

    meeting = next(item for item in result["meeting_projects"] if item["title"] == "新增自动识别项目")
    assert meeting["next"] == "安排上线"

    document_path = tmp_path / "部门周例会.docx"
    build_weekly_meeting_document(result, document_path)
    document_text = "\n".join(paragraph.text for table in Document(document_path).tables for row in table.rows for cell in row.cells for paragraph in cell.paragraphs)
    assert "新增自动识别项目（汇报人：张三）" in document_text
    assert "已完成联调" in document_text
    assert "安排上线" in document_text
    issue_paragraphs = [
        paragraph for table in Document(document_path).tables for row in table.rows for cell in row.cells
        for paragraph in cell.paragraphs
        if paragraph.text.startswith("本周问题") or paragraph.text.endswith("风险项")
    ]
    assert len(issue_paragraphs) >= 2
    assert all(
        run._r.find(qn("w:rPr")).find(qn("w:color")).get(qn("w:val")) == "FF0000"
        for paragraph in issue_paragraphs for run in paragraph.runs
    )
